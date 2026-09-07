"""
documents.py - Document upload, validation, parsing, and management router.

This module provides endpoints for:
- POST /documents/upload: Accepts multipart file uploads (PDF, DOCX, TXT, CSV), validates
  ownership, file format, and size (<=50MB), parses the text content using dedicated
  libraries (pdfplumber, python-docx, pandas), and stores the result in SQLite.
- GET /documents/{username}: Returns a lightweight list of document metadata for a user.
- DELETE /documents/{document_id}: Deletes a document if the requesting user owns it.
"""

import os
import tempfile
import logging
from typing import List, Optional
from datetime import datetime

import docx
import pandas as pd
import pdfplumber
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Document, User
from app.schemas import (
    DocumentMetadataResponse,
    DocumentUploadResponse,
    MessageResponse,
)
from app.rag import (
    chunk_text,
    chunk_csv_text,
    embed_and_store,
    delete_document_chunks,
    generate_summary,
)

logger = logging.getLogger(__name__)

# Initialize router under '/documents' prefix
router = APIRouter(prefix="/documents", tags=["Documents"])

# Supported file formats and size boundaries
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "csv"}
MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024  # 50 Megabytes


# ==============================================================================
# PARSING HELPER FUNCTIONS
# ==============================================================================

def parse_pdf_file(file_path: str) -> str:
    """
    Extracts plain text page-by-page from a PDF file using pdfplumber.
    Combines page texts with newlines.
    """
    pages_text = []
    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            extracted = page.extract_text()
            if extracted and extracted.strip():
                pages_text.append(extracted.strip())

    if not pages_text:
        raise ValueError("PDF file contains no readable text content (it may be scanned, image-only, or empty).")

    return "\n\n".join(pages_text)


def parse_docx_file(file_path: str) -> str:
    """
    Extracts all paragraph text from a Word DOCX file using python-docx.
    Combines paragraphs with newlines.
    """
    doc = docx.Document(file_path)
    paragraphs_text = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]

    # Also extract text from any tables in the document
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs_text.append(" | ".join(row_text))

    if not paragraphs_text:
        raise ValueError("DOCX document contains no readable text or paragraph content.")

    return "\n\n".join(paragraphs_text)


def parse_txt_file(file_path: str) -> str:
    """
    Reads plain text from a TXT file.
    Gracefully falls back to UTF-8 with errors='ignore' if encoding issues arise.
    """
    with open(file_path, "rb") as f:
        raw_bytes = f.read()

    try:
        text = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = raw_bytes.decode("utf-8", errors="ignore")

    clean_text = text.strip()
    if not clean_text:
        raise ValueError("TXT file is empty.")

    return clean_text


def parse_csv_file(file_path: str) -> str:
    """
    Parses a CSV file using pandas and converts each row into a structured readable string.
    Example line: "Row 1: column_a=val1, column_b=val2, ..."
    Includes column headers at the top to preserve table structure for RAG retrieval.

    NOTE FOR PROMPT 5:
    CSV data is converted to row-formatted plain text in this step. In Prompt 5 (Chunking & Embeddings),
    CSV documents will require row-aware chunking to prevent splitting individual records across chunks.
    """
    try:
        df = pd.read_csv(file_path)
    except Exception as exc:
        raise ValueError(f"Failed to read CSV file: {exc}")

    if df.empty:
        raise ValueError("CSV file contains no data rows.")

    columns = [str(col).strip() for col in df.columns]
    column_header_note = f"[Table Columns ({len(columns)})]: " + ", ".join(columns)

    formatted_rows = [column_header_note]
    for index, row in df.iterrows():
        row_fields = [f"{col}={row[col]}" for col in df.columns]
        formatted_rows.append(f"Row {index + 1}: " + ", ".join(row_fields))

    return "\n".join(formatted_rows)


def extract_document_text(file_path: str, file_type: str) -> str:
    """
    Dispatches file parsing to the appropriate helper based on file_type.
    """
    if file_type == "pdf":
        return parse_pdf_file(file_path)
    elif file_type == "docx":
        return parse_docx_file(file_path)
    elif file_type == "txt":
        return parse_txt_file(file_path)
    elif file_type == "csv":
        return parse_csv_file(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


# ==============================================================================
# ROUTE HANDLERS
# ==============================================================================

@router.post(
    "/upload",
    response_model=DocumentUploadResponse,
    status_code=status.HTTP_200_OK,
    summary="Upload and parse a document (PDF, DOCX, TXT, CSV)"
)
async def upload_document(
    username: str = Form(..., description="Registered username owning this document"),
    file: UploadFile = File(..., description="Document file to upload (max 50MB)"),
    db: Session = Depends(get_db),
):
    """
    Uploads and processes a document:
    1. Validates that the username exists in the users table.
    2. Validates file extension is one of (.pdf, .docx, .txt, .csv).
    3. Validates file size is <= 50MB while safely writing to temporary storage.
    4. Parses text based on file format.
    5. Stores document record with 'parsed' or 'failed' status and cleans up temp files.
    """
    # 1. Validate user existence
    user = db.query(User).filter(User.username == username).first()
    if not user:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"User '{username}' does not exist."
        )

    # 2. Validate file extension
    original_filename = file.filename or "unknown"
    extension = original_filename.split(".")[-1].lower() if "." in original_filename else ""

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file type. Please upload PDF, DOCX, TXT, or CSV."
        )

    # 3. Stream and write file to temporary file while checking size limit (50MB)
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=f".{extension}")
    temp_file_path = temp_file.name
    total_size = 0

    try:
        try:
            while chunk := await file.read(64 * 1024):  # Read 64KB chunks
                total_size += len(chunk)
                if total_size > MAX_FILE_SIZE_BYTES:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail="File size exceeds the 50MB limit."
                    )
                temp_file.write(chunk)
        finally:
            temp_file.close()

        # 4. Parse file contents into text
        parsed_text: Optional[str] = None
        doc_status = "parsed"
        error_msg: Optional[str] = None

        try:
            parsed_text = extract_document_text(temp_file_path, extension)
        except Exception as exc:
            logger.warning("Document parsing failed for '%s': %s", original_filename, exc)
            doc_status = "failed"
            error_msg = str(exc)

        # 5. Persist document record into SQLite database with initial status
        new_doc = Document(
            username=username,
            original_filename=original_filename,
            file_type=extension,
            extracted_text=parsed_text,
            status=doc_status,
            chunk_count=0,
            error_message=error_msg,
        )
        db.add(new_doc)
        db.commit()
        db.refresh(new_doc)

        # 6. Real-time RAG Ingestion: Chunk, embed, and store in ChromaDB
        if doc_status == "parsed" and parsed_text:
            try:
                # Step A: Chunk document content based on file type
                if extension == "csv":
                    # Row-aware CSV chunking with repeated column headers
                    chunks = chunk_csv_text(parsed_text, rows_per_chunk=20)
                else:
                    # Overlapping word-count chunking for PDF, DOCX, and TXT
                    chunks = chunk_text(parsed_text, chunk_size=500, overlap=50)

                if not chunks:
                    new_doc.status = "failed"
                    new_doc.chunk_count = 0
                    new_doc.error_message = "Document parsed but produced zero text chunks."
                    db.commit()
                    db.refresh(new_doc)

                    return DocumentUploadResponse(
                        id=new_doc.id,
                        filename=new_doc.original_filename,
                        status="failed",
                        chunk_count=0,
                        message="Document parsed but produced zero indexable text chunks.",
                        error_message="Document parsed but produced zero text chunks.",
                    )

                # Step B: Generate embeddings with sentence-transformers & store in ChromaDB
                stored_chunks = embed_and_store(
                    username=username,
                    document_id=new_doc.id,
                    filename=new_doc.original_filename,
                    file_type=new_doc.file_type,
                    chunks=chunks,
                )

                # Step C: Transition status to "indexed" upon successful vector storage
                new_doc.status = "indexed"
                new_doc.chunk_count = stored_chunks
                new_doc.error_message = None
                db.commit()
                db.refresh(new_doc)

                # Step D: Auto-generate document summary using Groq LLM (Prompt 8)
                # After reaching 'indexed' status, generate a concise summary from extracted_text.
                # If summary generation fails, leave summary as null and continue - never block or fail the response.
                try:
                    doc_summary = generate_summary(parsed_text, extension)
                    if doc_summary:
                        new_doc.summary = doc_summary
                        db.commit()
                        db.refresh(new_doc)
                except Exception as sum_err:
                    logger.error("Auto-summary generation failed for document %d: %s", new_doc.id, sum_err)
                    # Safe fallback: leave summary as None, don't fail upload

                return DocumentUploadResponse(
                    id=new_doc.id,
                    filename=new_doc.original_filename,
                    status="indexed",
                    chunk_count=new_doc.chunk_count,
                    summary=new_doc.summary,
                    message=f"Document uploaded and indexed successfully ({new_doc.chunk_count} chunks)",
                    error_message=None,
                )

            except Exception as index_exc:
                logger.error("Embedding and vector storage failed for doc %d: %s", new_doc.id, index_exc)
                new_doc.status = "failed"
                new_doc.chunk_count = 0
                new_doc.error_message = f"Vector indexing failed: {index_exc}"
                db.commit()
                db.refresh(new_doc)

                return DocumentUploadResponse(
                    id=new_doc.id,
                    filename=new_doc.original_filename,
                    status="failed",
                    chunk_count=0,
                    message=f"Document parsing succeeded, but vector indexing failed: {index_exc}",
                    error_message=f"Vector indexing failed: {index_exc}",
                )
        else:
            # Parsing stage itself failed
            return DocumentUploadResponse(
                id=new_doc.id,
                filename=new_doc.original_filename,
                status="failed",
                chunk_count=0,
                message=f"Document parsing failed: {error_msg}",
                error_message=error_msg,
            )

    finally:
        # Safely remove temporary file from disk
        if os.path.exists(temp_file_path):
            try:
                os.unlink(temp_file_path)
            except OSError as cleanup_err:
                logger.warning("Failed to clean up temp file '%s': %s", temp_file_path, cleanup_err)


@router.get(
    "/{username}",
    response_model=List[DocumentMetadataResponse],
    status_code=status.HTTP_200_OK,
    summary="List all uploaded documents metadata for a user"
)
def get_user_documents(
    username: str,
    db: Session = Depends(get_db),
):
    """
    Retrieves all document metadata (excluding full text content) for a given user.
    """
    user = db.query(User).filter(User.username == username).first()
    if not user:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"User '{username}' does not exist."
        )

    documents = (
        db.query(Document)
        .filter(Document.username == username)
        .order_by(Document.upload_timestamp.desc())
        .all()
    )

    return documents


@router.delete(
    "/{document_id}",
    response_model=MessageResponse,
    status_code=status.HTTP_200_OK,
    summary="Delete a document by ID with ownership verification"
)
def delete_document(
    document_id: int,
    username: str = Query(..., description="Username confirming document ownership"),
    db: Session = Depends(get_db),
):
    """
    Deletes an uploaded document:
    - Verifies document exists (404 if not).
    - Verifies requesting user is the owner (403 if not).
    - Removes record from database.
    """
    doc = db.query(Document).filter(Document.id == document_id).first()
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found."
        )

    if doc.username != username:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You do not have permission to delete this document."
        )

    # Purge associated vector chunks from ChromaDB to maintain consistency
    delete_document_chunks(username=username, document_id=document_id)

    db.delete(doc)
    db.commit()

    return MessageResponse(message="Document deleted successfully")
