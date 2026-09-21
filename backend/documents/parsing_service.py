"""
parsing_service.py - Document parsing and text extraction logic.

Handles text extraction for PDF, DOCX, TXT, and CSV file formats.
"""
from langfuse import observe


import docx
import pandas as pd
import pdfplumber

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "csv"}
MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024  # 50 Megabytes


def parse_pdf_file(file_path: str) -> str:
    """
    Extracts plain text page-by-page from a PDF file using pdfplumber.
    Combines page texts with newlines.
    """
    pages_text = []
    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            extracted = page.extract_text()
            if extracted and extracted.strip():
                pages_text.append(extracted.strip())

    if not pages_text:
        raise ValueError("PDF file contains no readable text content (it may be scanned, image-only, or empty).")

    return "\n\n".join(pages_text)


def parse_docx_file(file_path: str) -> str:
    """
    Extracts all paragraph text from a Word DOCX file using python-docx.
    Combines paragraphs with newlines.
    """
    doc = docx.Document(file_path)
    paragraphs_text = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]

    # Also extract text from any tables in the document
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs_text.append(" | ".join(row_text))

    if not paragraphs_text:
        raise ValueError("DOCX document contains no readable text or paragraph content.")

    return "\n\n".join(paragraphs_text)


def parse_txt_file(file_path: str) -> str:
    """
    Reads plain text from a TXT file.
    Gracefully falls back to UTF-8 with errors='ignore' if encoding issues arise.
    """
    with open(file_path, "rb") as f:
        raw_bytes = f.read()

    try:
        text = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = raw_bytes.decode("utf-8", errors="ignore")

    clean_text = text.strip()
    if not clean_text:
        raise ValueError("TXT file is empty.")

    return clean_text


def parse_csv_file(file_path: str) -> str:
    """
    Parses a CSV file using pandas and converts each row into a structured readable string.
    Example line: "Row 1: column_a=val1, column_b=val2, ..."
    Includes column headers at the top to preserve table structure for RAG retrieval.
    """
    try:
        df = pd.read_csv(file_path)
    except Exception as exc:
        raise ValueError(f"Failed to read CSV file: {exc}")

    if df.empty:
        raise ValueError("CSV file contains no data rows.")

    columns = [str(col).strip() for col in df.columns]
    column_header_note = f"[Table Columns ({len(columns)})]: " + ", ".join(columns)

    formatted_rows = [column_header_note]
    for index, row in df.iterrows():
        row_fields = [f"{col}={row[col]}" for col in df.columns]
        formatted_rows.append(f"Row {index + 1}: " + ", ".join(row_fields))

    return "\n".join(formatted_rows)


@observe(name="extract-text")
def extract_document_text(file_path: str, file_type: str) -> str:
    """
    Dispatches file parsing to the appropriate helper based on file_type.
    """
    if file_type == "pdf":
        return parse_pdf_file(file_path)
    elif file_type == "docx":
        return parse_docx_file(file_path)
    elif file_type == "txt":
        return parse_txt_file(file_path)
    elif file_type == "csv":
        return parse_csv_file(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


__all__ = [
    "ALLOWED_EXTENSIONS",
    "MAX_FILE_SIZE_BYTES",
    "parse_pdf_file",
    "parse_docx_file",
    "parse_txt_file",
    "parse_csv_file",
    "extract_document_text",
]
